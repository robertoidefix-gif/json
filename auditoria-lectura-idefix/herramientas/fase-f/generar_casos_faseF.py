"""Casos de borde de la fase F, con su verdad (lo que debe hacer Idefix1.0 con cada uno).

Uso: python3 generar_casos_faseF.py <carpeta_salida>   → <carpeta>/corpus/* y <carpeta>/verdad.json
Python solo GENERA los documentos; Idefix los lee tal cual en el navegador.
"""
import sys, os, io, json, re, zipfile
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib import pdfencrypt
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import docx
import openpyxl
import xlwt

OUT = sys.argv[1]
C = os.path.join(OUT, "corpus")
os.makedirs(C, exist_ok=True)
pdfmetrics.registerFont(TTFont("DejaVu", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))
VERDAD = []


def reg(i, archivo, nota, esperado, cajetilla="otros"):
    VERDAD.append({"id": i, "archivo": archivo, "cajetilla": cajetilla, "verdad": "faseF", "checks": [], "mime": "",
                   "nota": nota, "esperado": esperado})


def factura_pdf(c, numero, total, y0=800, base=None, cuota=None):
    c.setFont("DejaVu", 11)
    lineas = [f"Número de factura: {numero}", "Fecha de expedición: 10/04/2025", "Emisor: Construcciones Muñoz Ibáñez, S.L.",
              "NIF emisor: B12345674"]
    if base is not None:
        lineas += [f"Base imponible: {base}", f"IVA 21 %: {cuota}"]
    lineas.append(f"Total factura: {total}")
    for k, t in enumerate(lineas):
        c.drawString(50, y0 - 16 * k, t)


def pdf_bytes(dibujar, **kw):
    b = io.BytesIO()
    c = canvas.Canvas(b, pagesize=A4, **kw)
    dibujar(c)
    c.save()
    return b.getvalue()


# ---------------------------------------------------------------- F1 y F2: PDF
base_pdf = pdf_bytes(lambda c: (factura_pdf(c, "A-2025/0901", "121,00 €", base="100,00 €", cuota="21,00 €"), c.showPage()))
previos = b"Content-Type: application/pdf\r\n" + b"X" * 67 + b"\r\n"
assert len(previos) == 100
open(os.path.join(C, "f2_pdf_100_bytes_previos.pdf"), "wb").write(previos + base_pdf)
reg("f2a", "f2_pdf_100_bytes_previos.pdf", "100 bytes (cabecera de correo) antes de %PDF-: se lee con aviso informativo",
    {"estado": ["complete"], "aviso": "PDF_HEADER_OFFSET", "detalles": {"PDF_HEADER_OFFSET": {"bytesPrevios": 100}},
     "texto_contiene": ["Número de factura: A-2025/0901"]})
open(os.path.join(C, "f2_pdf_1030_bytes_previos.pdf"), "wb").write(b"Y" * 1030 + base_pdf)
reg("f2b", "f2_pdf_1030_bytes_previos.pdf", "1.030 bytes antes de %PDF- (más de 1.024): no es un PDF admisible",
    {"estado": ["failed"], "aviso": "FORMAT_REJECTED", "mensaje_contiene": "FORMATO NO ADMITIDO", "sin_aviso": ["PDF_HEADER_OFFSET"]})
open(os.path.join(C, "f2_pdf_basura.pdf"), "wb").write(b"\r\n" * 10 + b"%PDF-1.7\n" + bytes(range(256)) * 8)
reg("f2c", "f2_pdf_basura.pdf", "bytes previos y después %PDF- seguido de basura: rechazo controlado, sin texto técnico en inglés",
    {"estado": ["failed"], "mensaje_no_contiene": ["Invalid PDF structure", "No password given"]})
open(os.path.join(C, "f2_html_con_pdf.html"), "w", encoding="utf-8").write(
    "<!DOCTYPE html><html><head><meta charset=\"utf-8\"><title>Cita</title></head><body><p>La firma %PDF-1.4 aparece "
    "como texto en esta página.</p><p>Total factura: 50,00 €</p></body></html>")
reg("f2d", "f2_html_con_pdf.html", "página HTML que contiene «%PDF-» en sus primeros bytes: sigue siendo HTML",
    {"estado": ["complete"], "sin_aviso": ["PDF_HEADER_OFFSET"], "formato": "html", "texto_contiene": ["La firma %PDF-1.4 aparece"]})
cifrado = pdfencrypt.StandardEncryption("", ownerPassword="propietario", canPrint=0, canModify=0, canCopy=0)
open(os.path.join(C, "f1_pdf_solo_propietario.pdf"), "wb").write(
    pdf_bytes(lambda c: (factura_pdf(c, "A-2025/0902", "242,00 €", base="200,00 €", cuota="42,00 €"), c.showPage()), encrypt=cifrado))
reg("f1a", "f1_pdf_solo_propietario.pdf", "PDF con contraseña SOLO de propietario (restricciones de impresión y copia): se abre sin contraseña",
    {"estado": ["complete"], "sin_aviso": ["FORMAT_REJECTED"], "texto_contiene": ["Número de factura: A-2025/0902"]})
open(os.path.join(C, "f1_pdf_contrasena.pdf"), "wb").write(
    pdf_bytes(lambda c: (factura_pdf(c, "A-2025/0903", "1,00 €"), c.showPage()),
              encrypt=pdfencrypt.StandardEncryption("usuario", ownerPassword="propietario")))
reg("f1b", "f1_pdf_contrasena.pdf", "PDF con contraseña de apertura: rechazo en español, el texto técnico va en los detalles",
    {"estado": ["failed"], "aviso": "FORMAT_REJECTED", "mensaje_contiene": "PDF PROTEGIDO CON CONTRASEÑA",
     "mensaje_no_contiene": ["No password given"], "detalles": {"FORMAT_REJECTED": {"motivoTecnico": "PasswordException: No password given"}}})

# ---------------------------------------------------------------- F3 a F5: Word
W = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
     'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
     'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
     'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
     'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
     'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml"')


def p(texto, ppr=""):
    return f'<w:p>{ppr}<w:r><w:t xml:space="preserve">{texto}</w:t></w:r></w:p>'


def cuadro(*ramas):
    """Cuadro de texto: cada rama es ("Choice" | "Fallback", texto)."""
    partes = []
    for tipo, texto in ramas:
        cuerpo = f'<w:txbxContent>{p(texto)}</w:txbxContent>'
        if tipo == "Choice":
            partes.append(f'<mc:Choice Requires="wps"><w:drawing><wp:anchor><wp:docPr id="1" name="Cuadro"/><a:graphic><a:graphicData '
                          f'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp><wps:txbx>{cuerpo}'
                          f'</wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>')
        else:
            partes.append(f'<mc:Fallback><w:pict><v:shape><v:textbox>{cuerpo}</v:textbox></v:shape></w:pict></mc:Fallback>')
    return f'<w:p><w:r><mc:AlternateContent>{"".join(partes)}</mc:AlternateContent></w:r></w:p>'


def docx_con(nombre, cuerpo, numeracion=None, cabecera=None):
    d = docx.Document()
    d.add_paragraph("x")
    d.sections[0].header.paragraphs[0].text = "cabecera"
    b = io.BytesIO()
    d.save(b)
    zin = zipfile.ZipFile(io.BytesIO(b.getvalue()))
    original = zin.read("word/document.xml").decode("utf-8")
    sect = re.search(r"<w:sectPr\b.*?</w:sectPr>", original, re.S).group(0)
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as z:
        for info in zin.infolist():
            datos = zin.read(info.filename)
            if info.filename == "word/document.xml":
                datos = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {W}><w:body>{cuerpo}{sect}</w:body></w:document>'.encode()
            elif info.filename == "word/numbering.xml" and numeracion:
                datos = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:numbering {W}>{numeracion}</w:numbering>'.encode()
            elif info.filename == "word/header1.xml" and cabecera:
                datos = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:hdr {W}>{cabecera}</w:hdr>'.encode()
            z.writestr(info, datos)
    open(os.path.join(C, nombre), "wb").write(salida.getvalue())


docx_con("f3_docx_solo_fallback.docx", p("Antes del cuadro.") + cuadro(("Fallback", "Importe del cuadro VML: 300,00 €")) + p("Después."))
reg("f3a", "f3_docx_solo_fallback.docx", "cuadro de texto solo con mc:Fallback (VML): se lee una vez",
    {"estado": ["complete"], "texto_veces": [["Importe del cuadro VML: 300,00 €", 1]]})
docx_con("f3_docx_dos_choice.docx", cuadro(("Choice", "Opción A del cuadro"), ("Choice", "Opción B del cuadro"), ("Fallback", "Opción C del cuadro")))
reg("f3b", "f3_docx_dos_choice.docx", "dos mc:Choice y un mc:Fallback: solo cuenta el primer Choice",
    {"estado": ["complete"], "texto_veces": [["Opción A del cuadro", 1], ["Opción B del cuadro", 0], ["Opción C del cuadro", 0]]})
docx_con("f3_docx_cuadro_en_cabecera.docx", p("Cuerpo del documento."),
         cabecera=cuadro(("Choice", "Referencia de cabecera: EXP-77"), ("Fallback", "Referencia de cabecera: EXP-77")))
reg("f3c", "f3_docx_cuadro_en_cabecera.docx", "cuadro de texto (Choice y Fallback) en el encabezado: se lee una vez",
    {"estado": ["complete"], "texto_veces": [["Referencia de cabecera: EXP-77", 1]]})


def lvl(ilvl, fmt, texto, inicio=1, suff=None):
    return (f'<w:lvl w:ilvl="{ilvl}"><w:start w:val="{inicio}"/><w:numFmt w:val="{fmt}"/>' +
            (f'<w:suff w:val="{suff}"/>' if suff else "") + f'<w:lvlText w:val="{texto}"/></w:lvl>')


NUM = ("".join([
    '<w:abstractNum w:abstractNumId="10">' + lvl(0, "decimal", "%1.") + lvl(1, "decimal", "%1.%2.") + lvl(2, "lowerLetter", "%3)") + '</w:abstractNum>',
    '<w:abstractNum w:abstractNumId="11">' + lvl(0, "upperRoman", "%1.") + '</w:abstractNum>',
    '<w:abstractNum w:abstractNumId="12">' + lvl(0, "decimal", "Artículo %1.-", suff="space") + '</w:abstractNum>',
    '<w:abstractNum w:abstractNumId="13">' + lvl(0, "bullet", "\uf0b7") + '</w:abstractNum>',
    '<w:abstractNum w:abstractNumId="14">' + lvl(0, "decimalZero", "%1.") + '</w:abstractNum>',
    '<w:num w:numId="1"><w:abstractNumId w:val="10"/></w:num>',
    '<w:num w:numId="2"><w:abstractNumId w:val="10"/></w:num>',
    '<w:num w:numId="3"><w:abstractNumId w:val="10"/><w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>',
    '<w:num w:numId="4"><w:abstractNumId w:val="11"/></w:num>',
    '<w:num w:numId="5"><w:abstractNumId w:val="12"/></w:num>',
    '<w:num w:numId="6"><w:abstractNumId w:val="13"/></w:num>',
    '<w:num w:numId="7"><w:abstractNumId w:val="14"/></w:num>']))


def np_(num, ilvl=0):
    return f'<w:pPr><w:numPr><w:ilvl w:val="{ilvl}"/><w:numId w:val="{num}"/></w:numPr></w:pPr>'


celda = '<w:tbl><w:tr><w:tc>' + p("Romano en celda", np_(4)) + '</w:tc><w:tc>' + p("valor") + '</w:tc></w:tr></w:tbl>'
docx_con("f4_docx_numeracion.docx", "".join([
    p("Primero", np_(1)), p("Sub", np_(1, 1)), p("letra", np_(1, 2)), p("letra dos", np_(1, 2)), p("Sub dos", np_(1, 1)),
    p("Segundo", np_(1)), p("Tercero con otro numId", np_(2)), p("Reinicio", np_(3)),
    p("Romano uno", np_(4)), p("Romano dos", np_(4)), p("Objeto del contrato", np_(5)), p("Viñeta", np_(6)),
    p("Con cero", np_(7)), p("Sin número", '<w:pPr><w:numPr><w:numId w:val="0"/></w:numPr></w:pPr>'), celda]), numeracion=NUM)
reg("f4", "f4_docx_numeracion.docx", "numeración de varios niveles, continuación entre numId, reinicio, romanos, «Artículo %1.-», viñeta, "
    "decimalZero, numId 0 y lista en una celda", {"estado": ["complete"], "aviso": "DOCX_LIST_NUMBERING", "bloques_exactos": [
        "1. Primero", "1.1. Sub", "a) letra", "b) letra dos", "1.2. Sub dos", "2. Segundo", "3. Tercero con otro numId", "1. Reinicio",
        "I. Romano uno", "II. Romano dos", "Artículo 1.- Objeto del contrato", "• Viñeta", "01. Con cero", "Sin número"],
        "texto_contiene": ["III. Romano en celda"]})

R = '<w:r><w:t xml:space="preserve">{}</w:t></w:r>'
docx_con("f5_docx_revisiones.docx", "".join([
    '<w:p>' + R.format("Importe: ") + '<w:del w:id="1" w:author="R"><w:r><w:delText>900,00 € (ANTIGUO)</w:delText></w:r></w:del>' +
    '<w:ins w:id="2" w:author="R">' + R.format("950,00 €") + '</w:ins></w:p>',
    '<w:p><w:moveFrom w:id="3" w:author="R">' + R.format("Cláusula movida de sitio.") + '</w:moveFrom></w:p>',
    p("Texto intermedio."),
    '<w:p><w:moveTo w:id="4" w:author="R">' + R.format("Cláusula movida de sitio.") + '</w:moveTo></w:p>',
    '<w:p><w:r><w:rPr><w:b/><w:rPrChange w:id="5" w:author="R"><w:rPr/></w:rPrChange></w:rPr><w:t>Negrita nueva.</w:t></w:r></w:p>']))
reg("f5", "f5_docx_revisiones.docx", "inserción, eliminación, movimiento (moveFrom/moveTo) y cambio de formato sin aceptar",
    {"estado": ["partial"], "aviso": "DOCX_TRACKED_CHANGES",
     "detalles": {"DOCX_TRACKED_CHANGES": {"inserciones": 1, "eliminaciones": 1, "movimientos": 1, "cambiosFormato": 1}},
     "texto_veces": [["Cláusula movida de sitio.", 1], ["ANTIGUO", 0]], "texto_contiene": ["Importe: 950,00 €"]})

# ---------------------------------------------------------------- F6: Excel con celdas combinadas
CAB = ["Fecha expedición", "Serie", "Número factura", "NIF cliente", "Nombre cliente", "Base imponible", "Tipo IVA", "Cuota IVA", "Total", "Nota"]
FILAS = [["2025-01-0%d" % (k + 1), "A" if k < 3 else "B", "2025/%04d" % (k + 1), "B87654323", "Cliente %d" % (k + 1), 100.0, 21, 21.0, 121.0]
         for k in range(6)]
wb = openpyxl.Workbook()
ws = wb.active
ws.append(CAB)
for f in FILAS:
    ws.append(f)
for k in (3, 4, 6, 7):
    ws.cell(row=k, column=2).value = None
ws.merge_cells("B2:B4")
ws.merge_cells("B5:B7")
ws["J2"] = "revisada"
ws.merge_cells("J2:J1048576")      # rango enorme: debe expandirse solo en las filas con datos, sin bloquear
ws.merge_cells("K1:K30000")        # rango con la primera celda vacía: no se repite nada
ws["A10"] = "Total del libro"
ws.merge_cells("A10:E10")          # combinada en horizontal: las demás columnas quedan vacías
wb.save(os.path.join(C, "f6_xlsx_combinadas.xlsx"))
reg("f6a", "f6_xlsx_combinadas.xlsx", "Serie combinada en vertical (B2:B4, B5:B7), J2:J1048576, rango con origen vacío y combinada horizontal",
    {"estado": ["complete"], "aviso": "XLSX_MERGED_CELLS_EXPANDED", "combinadas": 4 + 6, "tiempo_max_ms": 15000,
     "celdas": [["2025/0003", "Serie", "A"], ["2025/0006", "Serie", "B"], ["2025/0004", "Nota", "revisada"]],
     "detalles": {"XLSX_MERGED_CELLS_EXPANDED": {"ranges": 5, "repeatedCells": 10}}}, cajetilla="libro_facturas_emitidas")
wx = xlwt.Workbook()
hx = wx.add_sheet("Libro")
for j, t in enumerate(CAB[:9]):
    hx.write(0, j, t)
for i, f in enumerate(FILAS):
    for j, v in enumerate(f):
        if j == 1:
            continue
        hx.write(i + 1, j, v)
hx.write_merge(1, 3, 1, 1, "A")
hx.write_merge(4, 6, 1, 1, "B")
wx.save(os.path.join(C, "f6_xls_combinadas.xls"))
reg("f6b", "f6_xls_combinadas.xls", "el mismo libro en .xls (registro MERGECELLS de BIFF8)",
    {"estado": ["complete"], "aviso": "XLSX_MERGED_CELLS_EXPANDED", "combinadas": 4,
     "celdas": [["2025/0003", "Serie", "A"], ["2025/0006", "Serie", "B"]]}, cajetilla="libro_facturas_emitidas")

# ---------------------------------------------------------------- F7: celdas de varias líneas en una tabla PDF
def tabla_varias_lineas(c):
    c.setFont("DejaVu", 10)
    xs = [50, 260, 340, 430]
    y = 780
    for x, t in zip(xs, ["Concepto", "Cantidad", "Precio", "Importe"]):
        c.drawString(x, y, t)
    filas = [(["Mantenimiento anual de", "la instalación eléctrica", "del local comercial"], "1", "400,00", "400,00"),
             (["Servicio de", "Limpieza Integral"], "2", "50,00", "100,00"),
             (["Retirada de", "escombros"], "1", "120,00", "120,00")]
    for conceptos, cant, precio, imp in filas:
        y -= 16
        c.drawString(xs[1], y, cant)
        c.drawString(xs[2], y, precio)
        c.drawString(xs[3], y, imp)
        for k, t in enumerate(conceptos):
            c.drawString(xs[0], y - 12 * k, t)
        y -= 12 * (len(conceptos) - 1)
    c.drawString(50, y - 12, "Observaciones: pago a 30 días")
    c.drawString(50, y - 40, "Base imponible: 620,00 €")
    c.showPage()


open(os.path.join(C, "f7_pdf_celdas_varias_lineas.pdf"), "wb").write(pdf_bytes(tabla_varias_lineas))
reg("f7", "f7_pdf_celdas_varias_lineas.pdf", "conceptos de 3 y 2 líneas (una continuación en mayúscula, entre filas) y una nota "
    "en mayúscula pegada debajo de la tabla, que no debe unirse",
    {"estado": ["complete"], "aviso": "PDF_TABLE_MULTILINE_CELLS",
     "celdas": [["400,00", "Concepto", "Mantenimiento anual de la instalación eléctrica del local comercial"],
                ["100,00", "Concepto", "Servicio de Limpieza Integral"], ["120,00", "Concepto", "Retirada de escombros"]],
     "ninguna_celda_contiene": ["Observaciones"]})

# ---------------------------------------------------------------- F8: varias facturas en un PDF
def dos_paginas_y_otra(c):
    factura_pdf(c, "A-2025/0950", "Suma y sigue")
    c.showPage()
    factura_pdf(c, "A-2025/0950", "605,00 €", base="500,00 €", cuota="105,00 €")
    c.showPage()
    factura_pdf(c, "A-2025/0951", "60,50 €", base="50,00 €", cuota="10,50 €")
    c.showPage()


open(os.path.join(C, "f8_pdf_factura_de_dos_paginas.pdf"), "wb").write(pdf_bytes(dos_paginas_y_otra))
reg("f8a", "f8_pdf_factura_de_dos_paginas.pdf", "una factura de dos páginas (número repetido en la 2.ª) y otra en la 3.ª: dos facturas",
    {"estado": ["complete"], "facturas": [["A-2025/0950", 1, 2, 605.0], ["A-2025/0951", 3, 3, 60.5]]})


def dos_en_una_pagina(c):
    factura_pdf(c, "A-2025/0960", "121,00 €", base="100,00 €", cuota="21,00 €")
    factura_pdf(c, "A-2025/0961", "242,00 €", y0=600, base="200,00 €", cuota="42,00 €")
    c.showPage()


open(os.path.join(C, "f8_pdf_dos_en_una_pagina.pdf"), "wb").write(pdf_bytes(dos_en_una_pagina))
reg("f8b", "f8_pdf_dos_en_una_pagina.pdf", "dos facturas en la misma página: no se parte (no se puede hacer con seguridad)",
    {"estado": ["complete"], "facturas": []})
open(os.path.join(C, "f8_pdf_una_factura.pdf"), "wb").write(base_pdf)
reg("f8c", "f8_pdf_una_factura.pdf", "una sola factura: documento.numero_documento como siempre, sin facturas_detectadas",
    {"estado": ["complete"], "facturas": [], "numero_documento": "A-2025/0901"})

json.dump(VERDAD, open(os.path.join(OUT, "verdad.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print(len(VERDAD), "casos en", OUT)
