"""Corpus de auditoría de Idefix1.0 — fase 1: fuentes HTML y documentos generados con Python/LibreOffice.

Uso: python3 generar_corpus.py <dir_salida>
La fase 2 (render en Chromium) la hace render_chromium.js y la fase 3 (derivados de imagen y PDF
escaneados) generar_corpus_imagenes.py.
"""
import sys, os, io, json, zipfile, shutil, subprocess, datetime, copy, html as htmlmod
from decimal import Decimal
from datos import *

OUT = os.path.abspath(sys.argv[1])
C = os.path.join(OUT, "corpus")
F = os.path.join(OUT, "fuentes")
T = os.path.join(OUT, "tmp")
for d in (C, F, T):
    os.makedirs(d, exist_ok=True)

VERDAD = []


def reg(id_, archivo, cajetilla, verdad, checks=None, mime="", nota=""):
    VERDAD.append({"id": id_, "archivo": archivo, "cajetilla": cajetilla, "verdad": verdad,
                   "checks": checks or [], "mime": mime, "nota": nota})


def soffice(args, cwd):
    r = subprocess.run(["soffice", "--headless", "--norestore"] + args, cwd=cwd, capture_output=True, text=True, timeout=180)
    if r.returncode != 0:
        raise RuntimeError(r.stderr + r.stdout)
    return r.stdout


# ============================ HTML fuente ============================
CSS = """body{font-family:'DejaVu Sans',Arial,sans-serif;font-size:14px;color:#111;background:#fff;margin:40px;width:714px}
h1{font-size:26px;margin:0 0 12px}p{margin:4px 0}table{border-collapse:collapse;margin:14px 0;width:100%}
th,td{border:1px solid #444;padding:5px 7px;text-align:left}td.n,th.n{text-align:right}.bloque{margin-top:12px}"""


def html_factura(extra_body="", head_extra="", charset="utf-8"):
    l = factura_lineas_texto()
    filas = "".join(f"<tr><td>{htmlmod.escape(c)}</td><td class=n>{q}</td><td class=n>{es(p)}</td><td class=n>{es(i)}</td></tr>"
                    for (c, q, p), i in zip(FACTURA["lineas"], FACTURA["lineas_importe"]))
    return f"""<!DOCTYPE html>
<html lang="es"><head><meta charset="{charset}"><title>Factura {FACTURA['numero']}</title><style>{CSS}</style>{head_extra}</head>
<body>
<h1>FACTURA</h1>
<p>{l[1]}</p>
<p>{l[2]}</p>
<div class="bloque"><p>{htmlmod.escape(l[3])}</p><p>{l[4]}</p><p>{htmlmod.escape(l[5])}</p></div>
<div class="bloque"><p>{htmlmod.escape(l[6])}</p><p>{l[7]}</p><p>{htmlmod.escape(l[8])}</p></div>
<table><thead><tr><th>Concepto</th><th class=n>Cantidad</th><th class=n>Precio</th><th class=n>Importe</th></tr></thead>
<tbody>{filas}</tbody></table>
<p>{l[-4]}</p>
<p>{l[-3]}</p>
<p><b>{l[-2]}</b></p>
<p>{l[-1]}</p>
{extra_body}
</body></html>
"""


def filas_libro_html(libro, tr_attr=""):
    out = []
    for r in libro:
        c = libro_fila_texto(r)
        out.append(f"<tr{tr_attr.format(n=r['numero'])}>" + "".join(
            f"<td{' class=n' if j >= 5 else ''}>{htmlmod.escape(v)}</td>" for j, v in enumerate(c)) + "</tr>")
    return "\n".join(out)


def html_libro(libro=LIBRO, titulo="Libro registro de facturas emitidas · Ejercicio 2025", tr_attr=""):
    cab = "".join(f"<th{' class=n' if j >= 5 else ''}>{h}</th>" for j, h in enumerate(LIBRO_CABECERA))
    return f"""<!DOCTYPE html>
<html lang="es"><head><meta charset="utf-8"><title>Libro de facturas emitidas</title><style>{CSS}
body{{width:auto;font-size:12px}} th,td{{padding:4px 5px}}</style></head>
<body>
<h1 style="font-size:20px">{titulo}</h1>
<p>Titular: {htmlmod.escape(EMISOR['nombre'])} · NIF: {EMISOR['nif']}</p>
<table><thead><tr>{cab}</tr></thead>
<tbody>
{filas_libro_html(libro, tr_attr)}
</tbody></table>
</body></html>
"""


# Libro largo (60 filas) para PDF multipágina
LIBRO60 = []
for k in range(60):
    base_r = LIBRO[k % 12]
    b = (base_r["base"] + Decimal(k * 7) + Decimal("0.35")).quantize(Decimal("0.01"))
    cuota = r2(b * base_r["tipo"] / 100)
    fecha = datetime.date(2025, 1 + k // 6, 1 + (k * 3) % 28)
    LIBRO60.append({**base_r, "fecha": fecha, "fecha_es": fecha.strftime("%d/%m/%Y"), "fecha_iso": fecha.isoformat(),
                    "numero": f"2025/{201 + k:04d}", "base": b, "cuota": cuota, "total": b + cuota})


def guardar(ruta, texto, cod="utf-8"):
    with open(ruta, "w", encoding=cod, newline="\n") as f:
        f.write(texto)


guardar(os.path.join(F, "factura.html"), html_factura())
guardar(os.path.join(F, "libro.html"), html_libro())
guardar(os.path.join(F, "libro_sin_bordes.html"), html_libro().replace("th,td{border:1px solid #444;", "th,td{border:none;"))
guardar(os.path.join(F, "libro60.html"), html_libro(LIBRO60).replace("<thead>", "<thead style=\"display:table-header-group\">"))

# ============================ HTML / HTM ============================
shutil.copy(os.path.join(F, "factura.html"), os.path.join(C, "html01_factura.html"))
reg("html01", "html01_factura.html", "otros", "factura", ["campos_factura", "tabla_lineas"], "text/html", "UTF-8, rótulos «Nombre: valor» y tabla de líneas")
shutil.copy(os.path.join(F, "libro.html"), os.path.join(C, "html02_libro_emitidas.html"))
reg("html02", "html02_libro_emitidas.html", "libro_facturas_emitidas", "libro", ["registros"], "text/html", "Tabla con thead/tbody")

lat = html_factura(charset="iso-8859-1")
with open(os.path.join(C, "htm03_latin1.htm"), "wb") as f:
    f.write(lat.encode("cp1252"))
reg("htm03", "htm03_latin1.htm", "otros", "factura", ["campos_factura", "sin_mojibake"], "text/html",
    "Declara iso-8859-1 y contiene € (0x80, windows-1252), como muchas páginas antiguas")
sin_meta = html_factura().replace('<meta charset="utf-8">', "")
with open(os.path.join(C, "htm04_cp1252_sin_meta.htm"), "wb") as f:
    f.write(sin_meta.encode("cp1252"))
reg("htm04", "htm04_cp1252_sin_meta.htm", "otros", "factura", ["campos_factura", "sin_mojibake"], "text/html",
    "Sin declaración de codificación, bytes windows-1252")

OCULTOS = {
    "OC_A": ('<p style="display:none">', "</p>", "display:none en línea"),
    "OC_B": ("<p hidden>", "</p>", "atributo hidden"),
    "OC_C": ('<p class="oculto">', "</p>", "clase CSS display:none en <style>"),
    "OC_D": ('<p style="color:#ffffff">', "</p>", "texto blanco sobre fondo blanco"),
    "OC_E": ('<p style="font-size:1px">', "</p>", "font-size:1px"),
    "OC_F": ('<p style="position:absolute;left:-9999px">', "</p>", "fuera de pantalla"),
    "OC_G": ('<div style="height:0;overflow:hidden" aria-hidden="true">', "</div>", "altura 0 + overflow hidden"),
    "OC_H": ('<p style="opacity:0">', "</p>", "opacity:0"),
}
extra = "\n".join(f"{a}{k}: {INYECCION}.{b}" for k, (a, b, _) in OCULTOS.items())
guardar(os.path.join(C, "html05_texto_oculto.html"), html_factura(extra, "<style>.oculto{display:none}</style>"))
reg("html05", "html05_texto_oculto.html", "otros", "factura", ["campos_factura", "ocultos_html"], "text/html",
    "8 variantes de texto oculto con una instrucción inyectada")

# rowspan: la serie ocupa una celda con rowspan=12 en la primera fila
cab = "".join(f"<th>{h}</th>" for h in LIBRO_CABECERA)
filas = []
for i, r in enumerate(LIBRO):
    c = libro_fila_texto(r)
    celdas = [f"<td>{htmlmod.escape(c[0])}</td>"]
    if i == 0:
        celdas.append(f'<td rowspan="{len(LIBRO)}">A</td>')
    celdas += [f"<td>{htmlmod.escape(v)}</td>" for v in c[2:]]
    filas.append("<tr>" + "".join(celdas) + "</tr>")
guardar(os.path.join(C, "html06_rowspan.html"), f"""<!DOCTYPE html><html lang="es"><head><meta charset="utf-8"><title>Libro</title></head><body>
<h1>Libro registro de facturas emitidas</h1><table><thead><tr>{cab}</tr></thead><tbody>
{chr(10).join(filas)}
</tbody></table></body></html>
""")
reg("html06", "html06_rowspan.html", "libro_facturas_emitidas", "libro", ["registros"], "text/html",
    "Columna Serie combinada verticalmente (rowspan=12), habitual en informes exportados")

guardar(os.path.join(C, "html07_filas_onclick.html"), html_libro(tr_attr=' onclick="verFactura(\'{n}\')"'))
reg("html07", "html07_filas_onclick.html", "libro_facturas_emitidas", "libro", ["registros"], "text/html",
    "Cada fila <tr> lleva onclick (habitual en páginas guardadas de portales bancarios o ERP web)")

frag = "<table>\n<tr>" + "".join(f"<th>{h}</th>" for h in LIBRO_CABECERA) + "</tr>\n" + filas_libro_html(LIBRO) + "\n</table>\n"
guardar(os.path.join(C, "html08_fragmento_tabla.html"), frag)
reg("html08", "html08_fragmento_tabla.html", "libro_facturas_emitidas", "libro", ["registros"], "text/html",
    "Fragmento HTML que empieza por <table> (sin <html> ni DOCTYPE), típico de exportaciones «Excel» de bancos")
guardar(os.path.join(C, "html09_meta_primero.html"),
        '<meta charset="utf-8">\n<title>Libro</title>\n<h1>Libro registro</h1>\n' + frag)
reg("html09", "html09_meta_primero.html", "libro_facturas_emitidas", "libro", ["registros"], "text/html",
    "Empieza por <meta charset> sin <html>")

ENT_HTML = ("<p>Importe: 1.234,56&nbsp;&euro; &mdash; Ca&ntilde;ada &amp; Hijos &#8364; &#x20AC; &hellip; "
            "&laquo;cita&raquo; &frac12; acci&oacute;n &ordm; &#128; &copy</p>")
ENT_ESPERADO = "Importe: 1.234,56 € — Cañada & Hijos € € … «cita» ½ acción º € ©"
guardar(os.path.join(C, "html11_entidades.html"),
        f'<!DOCTYPE html><html lang="es"><head><meta charset="utf-8"><title>Entidades</title></head><body>{ENT_HTML}</body></html>\n')
reg("html11", "html11_entidades.html", "otros", "texto", [{"tipo": "contiene", "texto": ENT_ESPERADO}], "text/html",
    "Entidades HTML con nombre, numéricas, 0x80 y sin punto y coma")

# ============================ DOCX (python-docx) ============================
import docx
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def docx_factura(totales_en_tabla=False):
    d = docx.Document()
    l = factura_lineas_texto()
    d.add_heading("FACTURA", level=1)
    for t in l[1:9]:
        d.add_paragraph(t)
    tb = d.add_table(rows=1, cols=4)
    tb.style = "Table Grid"
    for j, h in enumerate(FACTURA_TABLA[0]):
        tb.rows[0].cells[j].text = h
    for fila in FACTURA_TABLA[1:]:
        cs = tb.add_row().cells
        for j, v in enumerate(fila):
            cs[j].text = v
    if totales_en_tabla:
        t2 = d.add_table(rows=0, cols=2)
        t2.style = "Table Grid"
        for a, b in [("Base imponible", f"{es(FACTURA['base'])} €"), (f"IVA {FACTURA['tipo_iva']} %", f"{es(FACTURA['cuota'])} €"),
                     ("Total factura", f"{es(FACTURA['total'])} €")]:
            cs = t2.add_row().cells
            cs[0].text, cs[1].text = a, b
    else:
        for t in l[-4:-1]:
            d.add_paragraph(t)
    d.add_paragraph(l[-1])
    return d


docx_factura().save(os.path.join(C, "docx01_factura.docx"))
reg("docx01", "docx01_factura.docx", "otros", "factura", ["campos_factura", "tabla_lineas"], "", "python-docx: párrafos + tabla de líneas")
docx_factura(True).save(os.path.join(C, "docx02_totales_en_tabla.docx"))
reg("docx02", "docx02_totales_en_tabla.docx", "otros", "factura", ["campos_factura", "tabla_lineas"], "",
    "Base, IVA y total en una tabla de 2 columnas (maquetación muy habitual en facturas Word)")

d = docx.Document()
d.add_heading("Libro registro de facturas emitidas · Ejercicio 2025", level=1)
d.add_paragraph(f"Titular: {EMISOR['nombre']} · NIF: {EMISOR['nif']}")
tb = d.add_table(rows=1, cols=len(LIBRO_CABECERA))
tb.style = "Table Grid"
for j, h in enumerate(LIBRO_CABECERA):
    tb.rows[0].cells[j].text = h
for r in LIBRO:
    cs = tb.add_row().cells
    for j, v in enumerate(libro_fila_texto(r)):
        cs[j].text = v
d.save(os.path.join(C, "docx03_libro_emitidas.docx"))
reg("docx03", "docx03_libro_emitidas.docx", "libro_facturas_emitidas", "libro", ["registros"], "", "Tabla Word de 13×9")

# Listas numeradas y títulos
d = docx.Document()
d.add_heading("Documentación aportada", level=1)
LISTA_NUM = ["Libro registro de facturas emitidas", "Libro registro de facturas recibidas", "Modelo 303 del primer trimestre"]
for t in LISTA_NUM:
    d.add_paragraph(t, style="List Number")
d.add_heading("Observaciones", level=2)
LISTA_VIN = ["Falta el modelo 347", "Revisar la prorrata"]
for t in LISTA_VIN:
    d.add_paragraph(t, style="List Bullet")
d.save(os.path.join(C, "docx04_listas_titulos.docx"))
reg("docx04", "docx04_listas_titulos.docx", "otros", "texto",
    [{"tipo": "contiene", "texto": t} for t in LISTA_NUM + LISTA_VIN] +
    [{"tipo": "numeracion_lista", "items": [f"1. {LISTA_NUM[0]}", f"2. {LISTA_NUM[1]}", f"3. {LISTA_NUM[2]}"]},
     {"tipo": "titulos", "n": 2}], "", "Lista numerada automática (List Number), viñetas y títulos")


def docx_con_xml(nombre, fragmentos):
    """Crea un DOCX y sustituye párrafos marcadores @@k@@ por XML crudo."""
    d = docx.Document()
    d.add_paragraph("Informe de revisión de la factura A-2025/0157.")
    for k in fragmentos:
        d.add_paragraph(f"@@{k}@@")
    d.add_paragraph("Fin del informe.")
    b = io.BytesIO()
    d.save(b)
    zin = zipfile.ZipFile(io.BytesIO(b.getvalue()))
    zout_b = io.BytesIO()
    with zipfile.ZipFile(zout_b, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename == "word/document.xml":
                x = data.decode("utf-8")
                for k, xml in fragmentos.items():
                    p = f"<w:p><w:r><w:t>@@{k}@@</w:t></w:r></w:p>"
                    assert p in x, k
                    x = x.replace(p, xml)
                data = x.encode("utf-8")
            zout.writestr(it, data)
    with open(os.path.join(C, nombre), "wb") as f:
        f.write(zout_b.getvalue())


R = lambda t, rpr="": f'<w:r>{("<w:rPr>" + rpr + "</w:rPr>") if rpr else ""}<w:t xml:space="preserve">{t}</w:t></w:r>'
docx_con_xml("docx05_control_cambios.docx", {
    "cc": "<w:p>" + R("El importe correcto es ") +
          '<w:del w:id="1" w:author="Revisor" w:date="2025-03-20T10:00:00Z"><w:r><w:delText xml:space="preserve">1.200,00 € (IMPORTE ANTIGUO BORRADO)</w:delText></w:r></w:del>' +
          '<w:ins w:id="2" w:author="Revisor" w:date="2025-03-20T10:00:00Z">' + R("1.354,60 €") + "</w:ins>" +
          R(" según la factura rectificada.") + "</w:p>"})
reg("docx05", "docx05_control_cambios.docx", "otros", "texto",
    [{"tipo": "contiene", "texto": "El importe correcto es 1.354,60 € según la factura rectificada."},
     {"tipo": "no_contiene", "texto": "IMPORTE ANTIGUO BORRADO"},
     {"tipo": "aviso_esperado", "texto": "control de cambios", "patron": "(?i)revision|tracked|cambios|w:ins|w:del"}], "",
    "Control de cambios sin aceptar: una eliminación y una inserción")

docx_con_xml("docx06_texto_oculto.docx", {
    "oc": "<w:p>" + R("Texto visible del informe. ") + R(f"OCULTO_DOCX: {INYECCION}.", "<w:vanish/>") + "</w:p>"})
reg("docx06", "docx06_texto_oculto.docx", "otros", "texto",
    [{"tipo": "contiene", "texto": "Texto visible del informe."},
     {"tipo": "oculto_marcado", "marca": "OCULTO_DOCX"}], "",
    "Texto oculto de Word (w:vanish) con una instrucción inyectada")

TXBX = "Texto del cuadro: importe pendiente de cobro 1.000,00 €"
txbx_p = f"<w:p>{R(TXBX)}</w:p>"
docx_con_xml("docx07_cuadro_texto.docx", {"tb": (
    "<w:p><w:r><mc:AlternateContent><mc:Choice Requires=\"wps\"><w:drawing>"
    "<wp:anchor distT=\"0\" distB=\"0\" distL=\"114300\" distR=\"114300\" simplePos=\"0\" relativeHeight=\"1\" behindDoc=\"0\" locked=\"0\" layoutInCell=\"1\" allowOverlap=\"1\">"
    "<wp:simplePos x=\"0\" y=\"0\"/><wp:positionH relativeFrom=\"column\"><wp:posOffset>0</wp:posOffset></wp:positionH>"
    "<wp:positionV relativeFrom=\"paragraph\"><wp:posOffset>0</wp:posOffset></wp:positionV><wp:extent cx=\"3000000\" cy=\"600000\"/>"
    "<wp:effectExtent l=\"0\" t=\"0\" r=\"0\" b=\"0\"/><wp:wrapSquare wrapText=\"bothSides\"/><wp:docPr id=\"1\" name=\"Cuadro 1\"/><wp:cNvGraphicFramePr/>"
    "<a:graphic xmlns:a=\"http://schemas.openxmlformats.org/drawingml/2006/main\"><a:graphicData uri=\"http://schemas.microsoft.com/office/word/2010/wordprocessingShape\">"
    "<wps:wsp><wps:cNvSpPr txBox=\"1\"/><wps:spPr><a:xfrm><a:off x=\"0\" y=\"0\"/><a:ext cx=\"3000000\" cy=\"600000\"/></a:xfrm><a:prstGeom prst=\"rect\"><a:avLst/></a:prstGeom></wps:spPr>"
    f"<wps:txbx><w:txbxContent>{txbx_p}</w:txbxContent></wps:txbx><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>"
    "<mc:Fallback><w:pict><v:shape id=\"Cuadro1\" type=\"#_x0000_t202\" style=\"width:236pt;height:47pt\">"
    f"<v:textbox><w:txbxContent>{txbx_p}</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r></w:p>")})
reg("docx07", "docx07_cuadro_texto.docx", "otros", "texto",
    [{"tipo": "contiene", "texto": TXBX}, {"tipo": "apariciones", "texto": "importe pendiente de cobro", "n": 1}], "",
    "Cuadro de texto como lo guarda Word (mc:AlternateContent con Choice y Fallback VML)")

# Celdas combinadas
d = docx.Document()
d.add_paragraph("Resumen por cliente")
tb = d.add_table(rows=4, cols=4)
tb.style = "Table Grid"
datos_cc = [["Cliente", "Base", "Cuota", "Total"],
            ["Ferretería Núñez, S.L.", "1.250,00", "262,50", "1.512,50"],
            ["", "100,00", "21,00", "121,00"],
            ["Total general", "", "283,50", "1.633,50"]]
for i, f in enumerate(datos_cc):
    for j, v in enumerate(f):
        if v:
            tb.rows[i].cells[j].text = v
tb.cell(1, 0).merge(tb.cell(2, 0))   # vMerge
a = tb.cell(3, 0).merge(tb.cell(3, 1))  # gridSpan
d.save(os.path.join(C, "docx08_celdas_combinadas.docx"))
reg("docx08", "docx08_celdas_combinadas.docx", "otros", "tabla_combinada",
    [{"tipo": "celda_bajo_cabecera", "fila_contiene": "Total general", "columna": "Total", "valor": "1.633,50"},
     {"tipo": "celda_bajo_cabecera", "fila_contiene": "Total general", "columna": "Cuota", "valor": "283,50"},
     {"tipo": "celda_bajo_cabecera", "fila_contiene": "100,00", "columna": "Base", "valor": "100,00"}], "",
    "Celda combinada vertical (vMerge) y horizontal (gridSpan)")

d = docx.Document()
d.add_paragraph("Tabla con tabla anidada")
tb = d.add_table(rows=2, cols=2)
tb.style = "Table Grid"
tb.cell(0, 0).text = "Proveedor"
tb.cell(0, 1).text = "Detalle"
tb.cell(1, 0).text = "Ferretería Núñez, S.L."
inner = tb.cell(1, 1).add_table(rows=2, cols=2)
inner.cell(0, 0).text = "Base anidada"
inner.cell(0, 1).text = "7.777,77"
inner.cell(1, 0).text = "Cuota anidada"
inner.cell(1, 1).text = "1.633,33"
d.save(os.path.join(C, "docx09_tabla_anidada.docx"))
reg("docx09", "docx09_tabla_anidada.docx", "otros", "texto",
    [{"tipo": "contiene_en_algun_sitio", "texto": t} for t in ["7.777,77", "1.633,33", "Base anidada", "Ferretería Núñez, S.L."]], "",
    "Tabla dentro de una celda")

d = docx_factura()
sec = d.sections[0]
sec.header.paragraphs[0].text = f"ENCABEZADO: {EMISOR['nombre']} · Registro Mercantil de Madrid, tomo 123"
sec.footer.paragraphs[0].text = "PIE: Página 1 de 1 · Documento emitido electrónicamente"
d.save(os.path.join(C, "docx11_encabezado_pie.docx"))
reg("docx11", "docx11_encabezado_pie.docx", "otros", "factura",
    ["campos_factura", {"tipo": "contiene", "texto": "Registro Mercantil de Madrid, tomo 123"},
     {"tipo": "contiene", "texto": "PIE: Página 1 de 1"}], "", "Encabezado y pie de página")

b = open(os.path.join(C, "docx01_factura.docx"), "rb").read()
open(os.path.join(C, "docx12_truncado.docx"), "wb").write(b[: len(b) // 2])
reg("docx12", "docx12_truncado.docx", "otros", "error", [{"tipo": "rechazo_controlado"}], "", "ZIP truncado a la mitad")

# ============================ XLSX ============================
import openpyxl
from openpyxl.utils.datetime import CALENDAR_MAC_1904


def xlsx_libro(ruta, fila_cab=1, titulo=False, formulas=True, total=False, texto_es=False, epoch1904=False, combinar_serie=False):
    wb = openpyxl.Workbook()
    if epoch1904:
        wb.epoch = CALENDAR_MAC_1904
    ws = wb.active
    ws.title = "Emitidas"
    if titulo:
        ws.cell(1, 1, "Libro registro de facturas emitidas")
        ws.cell(2, 1, f"Ejercicio 2025 · Titular: {EMISOR['nombre']} · NIF: {EMISOR['nif']}")
    for j, h in enumerate(LIBRO_CABECERA, 1):
        ws.cell(fila_cab, j, h)
    for i, r in enumerate(LIBRO):
        f = fila_cab + 1 + i
        if texto_es:
            for j, v in enumerate(libro_fila_texto(r), 1):
                ws.cell(f, j, v)
            continue
        c = ws.cell(f, 1, r["fecha"]); c.number_format = "DD/MM/YYYY"
        ws.cell(f, 2, r["serie"]); ws.cell(f, 3, r["numero"]); ws.cell(f, 4, r["nif"]); ws.cell(f, 5, r["nombre"])
        c = ws.cell(f, 6, float(r["base"])); c.number_format = "#,##0.00"
        ws.cell(f, 7, r["tipo"])
        if formulas:
            c = ws.cell(f, 8, f"=ROUND(F{f}*G{f}/100,2)"); c.number_format = "#,##0.00"
            c = ws.cell(f, 9, f"=F{f}+H{f}"); c.number_format = "#,##0.00"
        else:
            c = ws.cell(f, 8, float(r["cuota"])); c.number_format = "#,##0.00"
            c = ws.cell(f, 9, float(r["total"])); c.number_format = "#,##0.00"
    if total:
        f = fila_cab + 1 + len(LIBRO)
        ws.cell(f, 1, "Total")
        for col in "FHI":
            c = ws[f"{col}{f}"]
            c.value = f"=SUM({col}{fila_cab + 1}:{col}{f - 1})"; c.number_format = "#,##0.00"
    if combinar_serie:
        ws.merge_cells(start_row=fila_cab + 1, start_column=2, end_row=fila_cab + len(LIBRO), end_column=2)
    wb.save(ruta)


def recalcular_lo(origen, destino_dir, fmt="xlsx", filtro=None):
    """Abre con LibreOffice Calc y guarda (calcula y guarda los valores de las fórmulas)."""
    arg = fmt if filtro is None else f"{fmt}:{filtro}"
    soffice(["--convert-to", arg, "--outdir", destino_dir, origen], T)


xlsx_libro(os.path.join(T, "xlsx01.xlsx"))
recalcular_lo(os.path.join(T, "xlsx01.xlsx"), T + "/lo", "xlsx", "Calc MS Excel 2007 XML")
shutil.copy(os.path.join(T, "lo", "xlsx01.xlsx"), os.path.join(C, "xlsx01_libro_emitidas.xlsx"))
reg("xlsx01", "xlsx01_libro_emitidas.xlsx", "libro_facturas_emitidas", "libro", ["registros", "valores_formula"], "",
    "Guardado por LibreOffice Calc: fechas reales, importes numéricos, cuota y total con fórmula (valor guardado)")

xlsx_libro(os.path.join(T, "xlsx02.xlsx"), fila_cab=4, titulo=True, total=True)
recalcular_lo(os.path.join(T, "xlsx02.xlsx"), T + "/lo", "xlsx", "Calc MS Excel 2007 XML")
shutil.copy(os.path.join(T, "lo", "xlsx02.xlsx"), os.path.join(C, "xlsx02_titulo_y_total.xlsx"))
reg("xlsx02", "xlsx02_titulo_y_total.xlsx", "libro_facturas_emitidas", "libro", ["registros", "fila_total_excluida"], "",
    "Dos filas de título, cabecera en la fila 4 y fila «Total» con SUMA")

xlsx_libro(os.path.join(C, "xlsx03_importes_texto_es.xlsx"), texto_es=True)
reg("xlsx03", "xlsx03_importes_texto_es.xlsx", "libro_facturas_emitidas", "libro", ["registros"], "",
    "Importes y fechas guardados como TEXTO en formato español («1.234,56», «14/03/2025»)")

xlsx_libro(os.path.join(C, "xlsx09_fechas_1904.xlsx"), formulas=False, epoch1904=True)
reg("xlsx09", "xlsx09_fechas_1904.xlsx", "libro_facturas_emitidas", "libro", ["registros"], "",
    "Sistema de fechas 1904 (libros creados en Excel para Mac)")

xlsx_libro(os.path.join(C, "xlsx07_celdas_combinadas.xlsx"), formulas=False, combinar_serie=True)
reg("xlsx07", "xlsx07_celdas_combinadas.xlsx", "libro_facturas_emitidas", "libro", ["registros"], "",
    "Columna Serie combinada verticalmente (B2:B13)")

# Varias hojas, una oculta con instrucción inyectada
xlsx_libro(os.path.join(T, "xlsx06.xlsx"), formulas=False)
wb = openpyxl.load_workbook(os.path.join(T, "xlsx06.xlsx"))
ws2 = wb.create_sheet("Notas")
ws2["A1"] = "Notas del cierre"; ws2["A2"] = "Revisado por el departamento contable"
ws3 = wb.create_sheet("Config")
ws3["A1"] = f"OCULTO_XLSX: {INYECCION}"
ws3.sheet_state = "hidden"
wb.save(os.path.join(C, "xlsx06_hoja_oculta.xlsx"))
reg("xlsx06", "xlsx06_hoja_oculta.xlsx", "libro_facturas_emitidas", "libro",
    ["registros", {"tipo": "hoja_oculta_marcada", "hoja": "Config"}, {"tipo": "inyeccion_detectada", "marca": "OCULTO_XLSX"}], "",
    "3 hojas; la hoja «Config» está oculta y contiene una instrucción inyectada")

# XLSX hechos a mano (XML) para fórmulas compartidas, celdas sin referencia r e inlineStr
CT_XLSX = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
           '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
           '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
           '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
           '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
           '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
RELS = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
WB = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Emitidas" sheetId="1" r:id="rId1"/></sheets></workbook>')
WBRELS = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
          '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
          '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
          '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/></Relationships>')
STYLES = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
          '<numFmts count="1"><numFmt numFmtId="164" formatCode="dd/mm/yyyy"/></numFmts>'
          '<fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts><fills count="1"><fill><patternFill patternType="none"/></fill></fills>'
          '<borders count="1"><border/></borders><cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
          '<cellXfs count="3"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/><xf numFmtId="164" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>'
          '<xf numFmtId="4" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/></cellXfs></styleSheet>')


def excel_serial(fecha):
    return (fecha - datetime.date(1899, 12, 30)).days


def xlsx_mano(nombre, modo):
    sst = []

    def s(t):
        if t not in sst:
            sst.append(t)
        return sst.index(t)

    filas = []
    col = "ABCDEFGHI"
    cab = []
    for j, h in enumerate(LIBRO_CABECERA):
        if modo == "inline":
            cab.append(f'<c r="{col[j]}1" t="inlineStr"><is><t>{htmlmod.escape(h)}</t></is></c>')
        elif modo == "sin_r":
            cab.append(f'<c t="s"><v>{s(h)}</v></c>')
        else:
            cab.append(f'<c r="{col[j]}1" t="s"><v>{s(h)}</v></c>')
    filas.append(f'<row r="1">{"".join(cab)}</row>' if modo != "sin_r" else f'<row>{"".join(cab)}</row>')
    for i, rr in enumerate(LIBRO):
        f = i + 2
        rf = (lambda c: f' r="{c}{f}"') if modo != "sin_r" else (lambda c: "")
        txt = lambda c, v: (f'<c{rf(c)} t="inlineStr"><is><t>{htmlmod.escape(v)}</t></is></c>' if modo == "inline"
                            else f'<c{rf(c)} t="s"><v>{s(v)}</v></c>')
        celdas = [f'<c{rf("A")} s="1"><v>{excel_serial(rr["fecha"])}</v></c>', txt("B", rr["serie"]), txt("C", rr["numero"]),
                  txt("D", rr["nif"]), txt("E", rr["nombre"]), f'<c{rf("F")} s="2"><v>{rr["base"]}</v></c>', f'<c{rf("G")}><v>{rr["tipo"]}</v></c>']
        if modo == "compartidas":
            if i == 0:
                fh = f'<f t="shared" ref="H2:H13" si="0">ROUND(F2*G2/100,2)</f>'
                fi = f'<f t="shared" ref="I2:I13" si="1">F2+H2</f>'
            else:
                fh, fi = '<f t="shared" si="0"/>', '<f t="shared" si="1"/>'
            celdas += [f'<c r="H{f}" s="2">{fh}<v>{rr["cuota"]}</v></c>', f'<c r="I{f}" s="2">{fi}<v>{rr["total"]}</v></c>']
        else:
            celdas += [f'<c{rf("H")} s="2"><v>{rr["cuota"]}</v></c>', f'<c{rf("I")} s="2"><v>{rr["total"]}</v></c>']
        filas.append(f'<row r="{f}">{"".join(celdas)}</row>' if modo != "sin_r" else f'<row>{"".join(celdas)}</row>')
    hoja = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            f'<sheetData>{"".join(filas)}</sheetData></worksheet>')
    sst_xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
               f'count="{len(sst)}" uniqueCount="{len(sst)}">' + "".join(f"<si><t>{htmlmod.escape(x)}</t></si>" for x in sst) + "</sst>")
    with zipfile.ZipFile(os.path.join(C, nombre), "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", CT_XLSX)
        z.writestr("_rels/.rels", RELS)
        z.writestr("xl/workbook.xml", WB)
        z.writestr("xl/_rels/workbook.xml.rels", WBRELS)
        z.writestr("xl/styles.xml", STYLES)
        z.writestr("xl/sharedStrings.xml", sst_xml)
        z.writestr("xl/worksheets/sheet1.xml", hoja)


xlsx_mano("xlsx04_formulas_compartidas.xlsx", "compartidas")
reg("xlsx04", "xlsx04_formulas_compartidas.xlsx", "libro_facturas_emitidas", "libro", ["registros", "valores_formula"], "",
    "Fórmulas compartidas (t=\"shared\") como las guarda Excel al «arrastrar» una fórmula")
xlsx_mano("xlsx05_celdas_sin_referencia.xlsx", "sin_r")
reg("xlsx05", "xlsx05_celdas_sin_referencia.xlsx", "libro_facturas_emitidas", "libro", ["registros"], "",
    "Celdas y filas sin atributo r (permitido por ECMA-376; lo hacen algunos generadores)")
xlsx_mano("xlsx08_cadenas_en_linea.xlsx", "inline")
reg("xlsx08", "xlsx08_cadenas_en_linea.xlsx", "libro_facturas_emitidas", "libro", ["registros"], "",
    "Cadenas en línea (inlineStr), habitual en exportaciones de ERP y bancos")

# Libro grande: 25.000 filas
wb = openpyxl.Workbook(write_only=True)
ws = wb.create_sheet("Emitidas")
ws.append(LIBRO_CABECERA)
for k in range(25000):
    r = LIBRO[k % 12]
    ws.append([r["fecha"], "B", f"2025/{k + 1:06d}", r["nif"], r["nombre"], float(r["base"]), r["tipo"], float(r["cuota"]), float(r["total"])])
wb.save(os.path.join(C, "xlsx10_grande_25000_filas.xlsx"))
reg("xlsx10", "xlsx10_grande_25000_filas.xlsx", "libro_facturas_emitidas", "grande",
    [{"tipo": "truncado_avisado", "filas": 25000}], "", "25.000 filas × 9 columnas (supera el límite de 20.000 filas)")

b = open(os.path.join(C, "xlsx01_libro_emitidas.xlsx"), "rb").read()
open(os.path.join(C, "xlsx11_truncado.xlsx"), "wb").write(b[: len(b) // 2])
reg("xlsx11", "xlsx11_truncado.xlsx", "libro_facturas_emitidas", "error", [{"tipo": "rechazo_controlado"}], "", "ZIP truncado a la mitad")

# ============================ XLS ============================
os.makedirs(T + "/xls", exist_ok=True)
for src, dst in [("lo/xlsx01.xlsx", "xls01_libro_emitidas.xls"), ("lo/xlsx02.xlsx", "xls02_titulo_y_total.xls")]:
    soffice(["--convert-to", "xls:MS Excel 97", "--outdir", T + "/xls", os.path.join(T, src)], T)
    shutil.copy(os.path.join(T, "xls", os.path.basename(src).replace(".xlsx", ".xls")), os.path.join(C, dst))
reg("xls01", "xls01_libro_emitidas.xls", "libro_facturas_emitidas", "libro", ["registros", "valores_formula"], "",
    "BIFF8 guardado por LibreOffice Calc (fechas, importes, fórmulas con valor guardado)")
reg("xls02", "xls02_titulo_y_total.xls", "libro_facturas_emitidas", "libro", ["registros", "fila_total_excluida"], "",
    "BIFF8: dos filas de título, cabecera en la fila 4 y fila «Total»")

shutil.copy(os.path.join(C, "xlsx03_importes_texto_es.xlsx"), os.path.join(T, "xlsx03.xlsx"))
soffice(["--convert-to", "xls:MS Excel 97", "--outdir", T + "/xls", os.path.join(T, "xlsx03.xlsx")], T)
shutil.copy(os.path.join(T, "xls", "xlsx03.xls"), os.path.join(C, "xls05_importes_texto_es.xls"))
reg("xls05", "xls05_importes_texto_es.xls", "libro_facturas_emitidas", "libro", ["registros"], "", "BIFF8 con importes y fechas como texto")

# SST con CONTINUE: 3000 cadenas únicas, mezcla de Latin-1 y caracteres de 16 bits
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Cadenas"
CADENAS = []
for k in range(3000):
    t = f"Proveedor {k:04d} · Peña Ibáñez · {'Łódź €' if k % 3 == 0 else 'Cañada'} · ref-{k * 7919 % 100000:05d}"
    CADENAS.append(t)
    ws.cell(k + 1, 1, t)
    ws.cell(k + 1, 2, k * 1.25)
wb.save(os.path.join(T, "sst.xlsx"))
soffice(["--convert-to", "xls:MS Excel 97", "--outdir", T + "/xls", os.path.join(T, "sst.xlsx")], T)
shutil.copy(os.path.join(T, "xls", "sst.xls"), os.path.join(C, "xls03_sst_continue.xls"))
json.dump(CADENAS, open(os.path.join(OUT, "cadenas_sst.json"), "w", encoding="utf-8"), ensure_ascii=False)
reg("xls03", "xls03_sst_continue.xls", "otros", "sst", [{"tipo": "cadenas_exactas", "n": 3000}], "",
    "3.000 cadenas únicas (tabla SST con registros CONTINUE, mezcla de 8 y 16 bits)")

import xlwt
wbx = xlwt.Workbook(encoding="utf-8")
sh = wbx.add_sheet("Emitidas")
st_fecha = xlwt.easyxf(num_format_str="DD/MM/YYYY")
st_imp = xlwt.easyxf(num_format_str="#,##0.00")
for j, h in enumerate(LIBRO_CABECERA):
    sh.write(0, j, h)
for i, r in enumerate(LIBRO, 1):
    sh.write(i, 0, r["fecha"], st_fecha)
    sh.write(i, 1, r["serie"]); sh.write(i, 2, r["numero"]); sh.write(i, 3, r["nif"]); sh.write(i, 4, r["nombre"])
    sh.write(i, 5, float(r["base"]), st_imp); sh.write(i, 6, r["tipo"])
    sh.write(i, 7, float(r["cuota"]), st_imp); sh.write(i, 8, float(r["total"]), st_imp)
wbx.save(os.path.join(C, "xls04_xlwt.xls"))
reg("xls04", "xls04_xlwt.xls", "libro_facturas_emitidas", "libro", ["registros"], "", "BIFF8 generado por otra librería (xlwt)")

b = open(os.path.join(C, "xls01_libro_emitidas.xls"), "rb").read()
open(os.path.join(C, "xls06_truncado.xls"), "wb").write(b[: len(b) // 2])
reg("xls06", "xls06_truncado.xls", "libro_facturas_emitidas", "error", [{"tipo": "rechazo_controlado"}], "", "OLE2 truncado a la mitad")

# ============================ Conversiones LibreOffice (realistas) ============================
# DOCX de Writer a partir de la factura HTML
os.makedirs(T + "/w", exist_ok=True)
soffice(["--infilter=HTML (StarWriter)", "--convert-to", "docx:MS Word 2007 XML", "--outdir", T + "/w", os.path.join(F, "factura.html")], T)
shutil.copy(os.path.join(T, "w", "factura.docx"), os.path.join(C, "docx10_writer.docx"))
reg("docx10", "docx10_writer.docx", "otros", "factura", ["campos_factura", "tabla_lineas"], "", "DOCX guardado por LibreOffice Writer")
# PDF de Writer a partir del DOCX
soffice(["--convert-to", "pdf", "--outdir", T + "/w", os.path.join(C, "docx01_factura.docx")], T)
shutil.copy(os.path.join(T, "w", "docx01_factura.pdf"), os.path.join(C, "pdf03_factura_writer.pdf"))
reg("pdf03", "pdf03_factura_writer.pdf", "otros", "factura", ["campos_factura", "tabla_lineas"], "", "PDF exportado por LibreOffice Writer desde Word")
# HTML exportado por Calc
soffice(["--convert-to", "html:HTML (StarCalc)", "--outdir", T + "/w", os.path.join(T, "lo", "xlsx01.xlsx")], T)
shutil.copy(os.path.join(T, "w", "xlsx01.html"), os.path.join(C, "htm10_calc_guardado_como_html.htm"))
reg("htm10", "htm10_calc_guardado_como_html.htm", "libro_facturas_emitidas", "libro", ["registros"], "text/html",
    "Hoja de cálculo «Guardar como página web» (LibreOffice Calc)")
# HTML exportado por Writer desde DOCX
soffice(["--convert-to", "html:HTML (StarWriter)", "--outdir", T + "/w", os.path.join(C, "docx01_factura.docx")], T)
shutil.copy(os.path.join(T, "w", "docx01_factura.html"), os.path.join(C, "html12_writer_guardado_como_html.html"))
reg("html12", "html12_writer_guardado_como_html.html", "otros", "factura", ["campos_factura", "tabla_lineas"], "text/html",
    "Documento Word «Guardar como página web» (LibreOffice Writer)")

# ============================ PDF (reportlab) ============================
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas as rlcanvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import pdfencrypt

FUENTE = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
pdfmetrics.registerFont(TTFont("DejaVu", FUENTE))
pdfmetrics.registerFont(TTFont("DejaVuB", FUENTE.replace("DejaVuSans.ttf", "DejaVuSans-Bold.ttf")))


def pdf_factura_canvas(c, y0=800, extra=None):
    l = factura_lineas_texto()
    y = y0
    c.setFont("DejaVuB", 20); c.drawString(50, y, "FACTURA"); y -= 30
    c.setFont("DejaVu", 10.5)
    for t in l[1:9]:
        c.drawString(50, y, t); y -= 16
    y -= 10
    xs = [50, 300, 380, 470]
    c.setFont("DejaVuB", 10.5)
    for x, h in zip(xs, FACTURA_TABLA[0]):
        c.drawString(x, y, h)
    y -= 16
    c.setFont("DejaVu", 10.5)
    for fila in FACTURA_TABLA[1:]:
        for j, (x, v) in enumerate(zip(xs, fila)):
            if j == 0:
                c.drawString(x, y, v)
            else:
                c.drawRightString(x + 60, y, v)
        y -= 16
    y -= 12
    for t in l[-4:]:
        c.drawString(50, y, t); y -= 16
    if extra:
        extra(c, y)
    return y


def pdf_simple(nombre, dibujar, encrypt=None, paginas=1):
    c = rlcanvas.Canvas(os.path.join(C, nombre), pagesize=A4, encrypt=encrypt)
    for p in range(paginas):
        dibujar(c, p)
        c.showPage()
    c.save()


pdf_simple("pdf01_factura_reportlab.pdf", lambda c, p: pdf_factura_canvas(c))
reg("pdf01", "pdf01_factura_reportlab.pdf", "otros", "factura", ["campos_factura", "tabla_lineas"], "", "PDF digital (reportlab)")


def dos_columnas(c, p):
    from reportlab.platypus import Frame, Paragraph
    from reportlab.lib.styles import ParagraphStyle
    st = ParagraphStyle("n", fontName="DejaVu", fontSize=10.5, leading=14, spaceAfter=10)
    c.setFont("DejaVuB", 16)
    c.drawString(50, 790, "Informe de revisión del IVA")
    f1 = Frame(50, 100, 235, 670, showBoundary=0)
    f2 = Frame(310, 100, 235, 670, showBoundary=0)
    f1.addFromList([Paragraph(t, st) for t in DOS_COLUMNAS_IZQ], c)
    f2.addFromList([Paragraph(t, st) for t in DOS_COLUMNAS_DER], c)


pdf_simple("pdf05_dos_columnas.pdf", dos_columnas)
reg("pdf05", "pdf05_dos_columnas.pdf", "otros", "dos_columnas", ["orden_lectura", {"tipo": "sin_tablas_falsas"}], "",
    "Texto maquetado en dos columnas (informes, resoluciones, BOE)")


def varias(c, p):
    c.setFont("DejaVuB", 20); c.drawString(50, 800, "FACTURA")
    c.setFont("DejaVu", 10.5)
    c.drawString(50, 770, f"Número de factura: A-2025/{301 + p:04d}")
    c.drawString(50, 754, f"Fecha de expedición: {10 + p:02d}/04/2025")
    c.drawString(50, 738, f"Emisor: {EMISOR['nombre']}")
    c.drawString(50, 722, f"NIF emisor: {EMISOR['nif']}")
    base = [Decimal("100.00"), Decimal("250.00"), Decimal("1000.00")][p]
    c.drawString(50, 690, f"Base imponible: {es(base)} €")
    c.drawString(50, 674, f"IVA 21 %: {es(r2(base * 21 / 100))} €")
    c.drawString(50, 658, f"Total factura: {es(base + r2(base * 21 / 100))} €")


pdf_simple("pdf07_tres_facturas.pdf", varias, paginas=3)
reg("pdf07", "pdf07_tres_facturas.pdf", "otros", "varias_facturas",
    [{"tipo": "varias_facturas", "numeros": ["A-2025/0301", "A-2025/0302", "A-2025/0303"], "totales": [121.0, 302.5, 1210.0]}], "",
    "Un PDF con tres facturas (una por página)")

# AcroForm
c = rlcanvas.Canvas(os.path.join(C, "pdf08_formulario_acroform.pdf"), pagesize=A4)
c.setFont("DejaVuB", 16); c.drawString(50, 800, "Solicitud de devolución")
c.setFont("DejaVu", 10.5)
form = c.acroForm
CAMPOS_FORM = {"nif": EMISOR["nif"], "nombre": "Construcciones Munoz Ibanez SL", "importe": "1.354,60", "iban": "ES9121000418450200051332"}
y = 760
for k, v in CAMPOS_FORM.items():
    c.drawString(50, y + 4, f"{k.upper()}:")
    form.textfield(name=k, value=v, x=150, y=y, width=300, height=18, borderWidth=1, fontName="Helvetica", fontSize=10)
    y -= 40
form.checkbox(name="conforme", x=150, y=y, size=14, checked=True, buttonStyle="check")
c.drawString(50, y + 3, "CONFORME:")
c.showPage(); c.save()
reg("pdf08", "pdf08_formulario_acroform.pdf", "otros", "texto",
    [{"tipo": "contiene_en_algun_sitio", "texto": v} for v in CAMPOS_FORM.values()], "",
    "Formulario PDF rellenable (AcroForm) con valores")

# Página mixta: texto digital (cabecera) + factura escaneada como imagen (se completa en fase 3)

# Texto oculto en PDF: blanco y modo de render 3 (invisible)
def oculto_pdf(c, p):
    def extra(c, y):
        c.setFillColorRGB(1, 1, 1)
        c.drawString(50, y - 20, f"OCULTO_PDF_BLANCO: {INYECCION}.")
        t = c.beginText(50, y - 40)
        t.setTextRenderMode(3)
        t.setFont("DejaVu", 10)
        t.textLine(f"OCULTO_PDF_INVISIBLE: {INYECCION}.")
        c.drawText(t)
        c.setFillColorRGB(0, 0, 0)
    pdf_factura_canvas(c, extra=extra)


pdf_simple("pdf10_texto_oculto.pdf", oculto_pdf)
reg("pdf10", "pdf10_texto_oculto.pdf", "otros", "factura",
    ["campos_factura", {"tipo": "inyeccion_detectada", "marca": "OCULTO_PDF_BLANCO"}, {"tipo": "inyeccion_detectada", "marca": "OCULTO_PDF_INVISIBLE"},
     {"tipo": "oculto_marcado", "marca": "OCULTO_PDF_INVISIBLE"}], "",
    "Instrucción inyectada en texto blanco y en texto invisible (modo de render 3)")

pdf_simple("pdf11a_cifrado_con_clave.pdf", lambda c, p: pdf_factura_canvas(c),
           encrypt=pdfencrypt.StandardEncryption("1234", ownerPassword="dueño", strength=128))
reg("pdf11a", "pdf11a_cifrado_con_clave.pdf", "otros", "error", [{"tipo": "rechazo_controlado", "patron": "(?i)contrase|cifrad|password|protegid"}], "",
    "PDF con contraseña de apertura")
pdf_simple("pdf11b_cifrado_solo_propietario.pdf", lambda c, p: pdf_factura_canvas(c),
           encrypt=pdfencrypt.StandardEncryption("", ownerPassword="dueño", canPrint=1, canModify=0, canCopy=0, strength=128))
reg("pdf11b", "pdf11b_cifrado_solo_propietario.pdf", "otros", "factura", ["campos_factura"], "",
    "PDF con restricciones de propietario y sin contraseña de apertura (habitual en extractos bancarios y certificados)")

b = open(os.path.join(C, "pdf01_factura_reportlab.pdf"), "rb").read()
open(os.path.join(C, "pdf12_truncado.pdf"), "wb").write(b[: len(b) // 2])
reg("pdf12", "pdf12_truncado.pdf", "otros", "error_o_parcial", [{"tipo": "rechazo_controlado"}], "", "PDF truncado a la mitad")


def rotado(c, p):
    c.setPageRotation(90)
    pdf_factura_canvas(c)


pdf_simple("pdf13_pagina_rotada.pdf", rotado)
reg("pdf13", "pdf13_pagina_rotada.pdf", "otros", "factura", ["campos_factura"], "", "Página con /Rotate 90")
open(os.path.join(C, "pdf14_bytes_previos.pdf"), "wb").write(b"\r\n" + b)
reg("pdf14", "pdf14_bytes_previos.pdf", "otros", "factura", ["campos_factura"], "",
    "Dos bytes (CRLF) antes de %PDF- (los visores admiten la cabecera en los primeros 1.024 bytes)")

json.dump({"LIBRO60": [{**r, "fecha": r["fecha_iso"], "base": str(r["base"]), "cuota": str(r["cuota"]), "total": str(r["total"])} for r in LIBRO60]},
          open(os.path.join(OUT, "libro60.json"), "w", encoding="utf-8"), ensure_ascii=False)
json.dump(VERDAD, open(os.path.join(OUT, "verdad_fase1.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print(len(VERDAD), "documentos en la fase 1")
